from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/saikarthik26/Documents/Codex/2026-09-03/https-www-sih-gov-in-sih2026psgo')
OUT = ROOT / 'outputs' / 'SIH26170_Teammate_One_Day_Learning_Sheet.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color='D9D9D9', size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:color'), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + name))
        if node is None:
            node = OxmlElement('w:' + name)
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths=None, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, '1F4E78')
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for r_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_border(cell)
            set_cell_margins(cell)
            if r_idx % 2:
                set_cell_shading(cell, 'F3F7FA')
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(text))
            run.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_lead(doc, lead, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(lead)
    r.bold = True
    p.add_run(body)
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.67)
sec.bottom_margin = Inches(0.67)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(32, 32, 32)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for name, size, before, after in [
    ('Title', 25, 0, 8),
    ('Subtitle', 12, 0, 15),
    ('Heading 1', 17, 13, 7),
    ('Heading 2', 12.5, 10, 5),
    ('Heading 3', 11, 8, 4),
]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name in ('Title', 'Heading 1') else 'Aptos'
    st.font.size = Pt(size)
    st.font.bold = name != 'Subtitle'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run('SIH26170 teammate learning sheet')
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph('SIH26170 Teammate Learning Sheet', style='Title')
sub = doc.add_paragraph('AI driven anomaly detection in component burn in and screening', style='Subtitle')
sub.alignment = WD_ALIGN_PARAGRAPH.LEFT

add_lead(doc, 'Purpose. ', 'Use this sheet to understand the problem in one day and begin building a defensible nomination prototype. Nobody needs to become a semiconductor expert before starting. The team must understand the measurements, the two required AI tasks, the limits of synthetic data, and how the demo will be evaluated.')
add_lead(doc, 'The main idea. ', 'Traditional testing waits for an electrical value to cross a fixed failure limit. Our system identifies components that behave unusually compared with their batch and predicts whether early drift will become dangerous later.')

doc.add_heading('The problem in sixty seconds', level=1)
doc.add_paragraph('Electronic components intended for high reliability applications are operated under controlled stress for many hours. Engineers periodically measure values such as leakage current, standby current or propagation delay. A component may remain inside its official limit while already drifting much faster than the rest of its batch. The goal is to catch that component early.')

add_table(doc, ['Traditional test', 'Our proposed system'], [
    ['Checks whether the current value has crossed a fixed limit.', 'Checks the fixed limit, batch-relative behaviour and change over time.'],
    ['Usually reacts after a clear abnormality appears.', 'Forecasts the later value from early measurements.'],
    ['Returns pass or fail.', 'Returns accept, monitor, retest or reject with reasons and uncertainty.'],
], widths=[3.45, 3.55], font_size=9.7)

doc.add_heading('A simple example', level=2)
doc.add_paragraph('Suppose the leakage-current limit is 50 microamps. Most devices in a batch measure about 10 microamps. One device measures 40 microamps at 24 hours. It technically passes, but it is far from the batch norm and may be rising quickly. The system should flag it before it reaches the limit.')

doc.add_heading('The two required AI tasks', level=1)
add_number(doc, 'Dynamic outlier detection: find components that are unusual relative to comparable devices, even when their values are still within absolute limits.')
add_number(doc, 'Drift prediction: use early measurements, such as 0-hour and 24-hour values, to predict the value at 168 hours and estimate failure risk.')
add_lead(doc, 'Optional extension. ', 'Rank probable degradation mechanisms such as thermal overstress or moisture-related damage. Call these probable causes unless controlled experiments or confirmed failure analysis provide true cause labels.')

doc.add_page_break()
doc.add_heading('What the final prototype should do', level=1)
add_bullet(doc, 'Upload a CSV containing component, batch, stress and time-series measurements.')
add_bullet(doc, 'Group components by comparable family and manufacturing batch.')
add_bullet(doc, 'Apply fixed safety limits and batch-relative anomaly detection.')
add_bullet(doc, 'Predict the 168-hour measurement from early observations.')
add_bullet(doc, 'Display a prediction interval or probability of exceeding the limit.')
add_bullet(doc, 'Explain the features that caused each warning.')
add_bullet(doc, 'Recommend accept, monitor, retest or reject.')
add_bullet(doc, 'Compare the intelligent system with a fixed-threshold baseline.')

doc.add_heading('Our proposed solution', level=1)
doc.add_paragraph('Use a hybrid approach. Simple physical and statistical features keep the system understandable; machine learning adds batch-aware anomaly detection and future-value prediction.')

flow = [
    ('1  Data input', 'Measurements at 0h, 24h, 96h and 168h; batch; component family; temperature; humidity; test condition.'),
    ('2  Data checks', 'Units, missing values, impossible readings, duplicates and component-history consistency.'),
    ('3  Feature engineering', 'Change, percentage change, drift slope, acceleration, batch deviation, variability and distance from the safety limit.'),
    ('4  Outlier model', 'Isolation Forest plus a robust statistical baseline such as median absolute deviation.'),
    ('5  Drift predictor', 'Linear regression baseline and XGBoost main model for predicting the 168-hour value.'),
    ('6  Uncertainty', 'Conformal prediction or quantile regression to produce a credible prediction range.'),
    ('7  Explanation', 'SHAP values and plain-language reason codes.'),
    ('8  Decision layer', 'Accept, monitor, retest or reject using asymmetric risk rules.'),
    ('9  Dashboard', 'Batch overview, component trajectory, prediction range, reason and recommended action.'),
]
add_table(doc, ['Stage', 'What happens'], flow, widths=[1.45, 5.55], font_size=9.5)

doc.add_heading('Why these models', level=2)
add_bullet(doc, 'Isolation Forest learns unusual combinations without requiring many labelled failures.')
add_bullet(doc, 'Robust statistics provide an interpretable benchmark and expose whether ML is genuinely helping.')
add_bullet(doc, 'XGBoost works well for small and medium tabular datasets and does not require long sequences.')
add_bullet(doc, 'Conformal prediction adds uncertainty without requiring a complicated Bayesian neural network.')
add_bullet(doc, 'SHAP explains which early measurements and drift features influenced the prediction.')

doc.add_page_break()
doc.add_heading('Do not start with an LSTM', level=2)
doc.add_paragraph('If each component has only four timestamps, deep sequence models are unnecessary. Start with linear regression and XGBoost. Test an LSTM only if the dataset later contains many closely spaced measurements per component.')

doc.add_heading('Decision logic', level=1)
add_table(doc, ['Result', 'Meaning', 'Action'], [
    ['Accept', 'Normal trajectory and low predicted risk.', 'Continue the approved process.'],
    ['Monitor', 'Small deviation but no strong failure evidence.', 'Watch future measurements.'],
    ['Retest', 'Prediction is uncertain or data quality is weak.', 'Collect another measurement or inspect setup.'],
    ['Reject early', 'High anomaly score and high predicted limit-crossing risk.', 'Send for engineer review and possible rejection.'],
], widths=[1.1, 3.65, 2.25], font_size=9.3)

doc.add_heading('One day team learning plan', level=1)
doc.add_paragraph('The objective is not mastery. By the end of the day, every teammate should be able to explain the problem, read the dataset, describe the pipeline and defend one part of the implementation.')

add_table(doc, ['Time', 'Learn', 'Practical output'], [
    ['09:00 to 09:45', 'Problem and vocabulary', 'Explain burn-in, drift, absolute limit, batch-relative anomaly and false negative.'],
    ['09:45 to 10:45', 'Data structure', 'Open a sample CSV and trace one component across time.'],
    ['11:00 to 12:00', 'Statistics and features', 'Calculate change, slope, batch median and median absolute deviation.'],
    ['12:00 to 13:00', 'Anomaly detection', 'Run or understand a robust threshold and Isolation Forest.'],
    ['14:00 to 15:00', 'Prediction', 'Train or understand linear regression and XGBoost for the 168-hour value.'],
    ['15:00 to 15:45', 'Evaluation', 'Explain false negatives, false positives, MAE, recall and calibration.'],
    ['15:45 to 16:30', 'Uncertainty and explanations', 'Read a prediction interval and explain one SHAP chart.'],
    ['16:30 to 17:30', 'Product and demo', 'Walk through upload, batch view, flagged component and recommendation.'],
    ['17:30 to 18:00', 'Team rehearsal', 'Each person gives a 60-second explanation of their module.'],
], widths=[1.25, 2.25, 3.5], font_size=9.0)

doc.add_heading('The minimum skills to learn', level=1)
add_table(doc, ['Skill', 'Minimum understanding required', 'Not required on day one'], [
    ['Semiconductor testing', 'Why components are stressed and why electrical values are measured over time.', 'Device physics, fabrication-process mastery or certification standards.'],
    ['Python and pandas', 'Read CSV, group by batch/component, calculate features and plot trajectories.', 'Advanced software architecture.'],
    ['Statistics', 'Median, deviation, slope, outlier and prediction interval.', 'Advanced probability proofs.'],
    ['Machine learning', 'Fit, predict, train/test split, overfitting and feature importance.', 'Designing a neural architecture.'],
    ['Evaluation', 'False negatives matter most; components and batches must not leak across splits.', 'Research-level statistical inference.'],
    ['Dashboard', 'Display evidence and decisions clearly.', 'Complex animations or production infrastructure.'],
], widths=[1.2, 3.65, 2.15], font_size=8.9)

doc.add_heading('Useful learning resources to search for', level=2)
add_bullet(doc, 'Pandas groupby and time-series feature engineering.')
add_bullet(doc, 'Median absolute deviation for robust outlier detection.')
add_bullet(doc, 'Scikit-learn Isolation Forest tutorial.')
add_bullet(doc, 'XGBoost regression tutorial.')
add_bullet(doc, 'SHAP feature importance tutorial.')
add_bullet(doc, 'Conformal prediction for regression using MAPIE.')

doc.add_page_break()
doc.add_heading('Team roles and realistic responsibilities', level=1)
doc.add_paragraph('For a four-person team, assign one owner per lane. Everyone should still understand the complete flow.')

add_table(doc, ['Owner', 'Primary responsibilities', 'End-of-day deliverable'], [
    ['Data and domain', 'Dataset schema, synthetic generator, batch logic, data checks and degradation-pattern documentation.', 'Clean CSV plus data dictionary and five plotted trajectories.'],
    ['Anomaly model', 'Robust statistical baseline, Isolation Forest, anomaly explanations and fault-injection tests.', 'Ranked suspicious components with baseline comparison.'],
    ['Prediction model', 'Feature engineering, linear baseline, XGBoost, uncertainty and leakage-safe evaluation.', '168-hour predictions with MAE and prediction intervals.'],
    ['Product and integration', 'FastAPI or direct Python integration, Streamlit dashboard, decision rules and presentation.', 'Working end-to-end demonstration.'],
], widths=[1.25, 3.95, 1.8], font_size=9.0)

doc.add_heading('Dataset design', level=1)
doc.add_paragraph('Use long-format data so every row represents one component at one time. This supports different measurement schedules and makes plots and validation easier.')

add_table(doc, ['Field', 'Example', 'Why it exists'], [
    ['component_id', 'C0024', 'Keeps all readings from one component together.'],
    ['batch_id', 'B07', 'Enables batch-relative comparison and leakage-safe splitting.'],
    ['component_family', 'Digital IC', 'Prevents incomparable component types from being mixed.'],
    ['hours', '0, 24, 96, 168', 'Defines the burn-in timeline.'],
    ['temperature_c', '125', 'Records applied thermal stress.'],
    ['humidity_pct', '20', 'Optional environmental exposure.'],
    ['leakage_ua', '17.4', 'Example degradation-sensitive measurement.'],
    ['propagation_delay_ns', '5.8', 'Optional secondary electrical measurement.'],
    ['cause_label', 'thermal stress', 'Only when simulated or confirmed by controlled evidence.'],
    ['final_result', 'fail', 'Training and evaluation target.'],
    ['data_source', 'synthetic', 'Separates simulated, public and laboratory records.'],
], widths=[1.55, 1.35, 4.1], font_size=8.8)

doc.add_heading('Synthetic-data rules', level=2)
add_bullet(doc, 'Generate the full trajectory for each component before splitting the data.')
add_bullet(doc, 'Include healthy, gradual, accelerating, sudden and intermittent patterns.')
add_bullet(doc, 'Add manufacturing variation, measurement noise and batch-to-batch differences.')
add_bullet(doc, 'Do not use the cause label to generate features that would reveal the answer directly.')
add_bullet(doc, 'Split entire batches between training, validation and testing.')
add_bullet(doc, 'Label synthetic records visibly and never claim aerospace validation from them.')

doc.add_heading('Probable degradation causes', level=2)
add_table(doc, ['Cause', 'Possible observable pattern', 'Required caution'], [
    ['Thermal overstress', 'Accelerating leakage or timing drift under elevated temperature.', 'Other faults may create similar patterns.'],
    ['Moisture ingress', 'Humidity exposure followed by leakage increase or instability.', 'Needs exposure history or failure analysis for confirmation.'],
    ['Electrical overstress', 'Sudden step change after an abnormal voltage/current event.', 'Requires event or test-condition records.'],
    ['Manufacturing variation', 'Stable but batch-relative offset from the start.', 'May be harmless if the trajectory remains stable.'],
    ['Measurement fault', 'Impossible spike, frozen reading or inconsistent channel behaviour.', 'Do not misclassify tester failure as component failure.'],
], widths=[1.35, 3.65, 2.0], font_size=8.8)

doc.add_page_break()
doc.add_heading('Evaluation that judges can trust', level=1)
add_lead(doc, 'Most important principle. ', 'A false negative means a defective component is accepted. The project must prioritize catching defective components while keeping unnecessary rejection visible.')

add_table(doc, ['Metric', 'Question it answers'], [
    ['Defect recall', 'Of all truly defective components, how many did we catch?'],
    ['False-negative rate', 'How many defective components escaped?'],
    ['False-positive rate', 'How many healthy components were unnecessarily flagged?'],
    ['168-hour MAE', 'How far were final-value predictions from the truth?'],
    ['Interval coverage', 'Did the claimed uncertainty ranges contain the true values often enough?'],
    ['Early-warning lead time', 'How much earlier was the problem identified?'],
    ['Cause top-k accuracy', 'Was the confirmed degradation cause among the highest-ranked possibilities?'],
], widths=[1.75, 5.25], font_size=9.3)

doc.add_heading('Required comparisons', level=2)
add_table(doc, ['Baseline', 'Purpose'], [
    ['Fixed datasheet limit', 'Represents the conventional pass/fail rule.'],
    ['Batch median plus robust deviation', 'Tests whether simple statistics are sufficient.'],
    ['Linear regression', 'Provides an interpretable future-value baseline.'],
    ['Isolation Forest plus XGBoost', 'Represents the proposed intelligent pipeline.'],
], widths=[2.6, 4.4], font_size=9.4)

doc.add_heading('Data leakage to avoid', level=2)
add_bullet(doc, 'Never put early readings from one component in training and its later readings in testing.')
add_bullet(doc, 'Do not randomly split individual rows. Split by component and preferably by batch.')
add_bullet(doc, 'Do not use the 96-hour or 168-hour value to construct a feature when claiming prediction from 0h and 24h.')
add_bullet(doc, 'Fit normalizers and thresholds using training data only.')

doc.add_heading('The nomination demonstration', level=1)
for step_number, step_text in enumerate([
    'Show a batch in which every component is still below the absolute limit.',
    'Select one component whose early trajectory is abnormal.',
    'Show why the fixed threshold passes it.',
    'Show the model forecasting that it may cross the limit by 168 hours.',
    'Display the prediction range, batch comparison and most important features.',
    'Apply the recommendation: monitor, retest or reject early.',
    'Reveal the simulated final reading and compare both approaches.',
], start=1):
    doc.add_paragraph(f'{step_number}.  {step_text}')

doc.add_heading('The thirty second team pitch', level=2)
doc.add_paragraph('Current burn-in screening mainly catches components after a parameter crosses a fixed limit. Our system also learns what is normal for each component batch and studies how individual measurements change over time. Using early readings, it predicts the final test value, estimates uncertainty and explains whether a component should be accepted, monitored, retested or rejected. The aim is to identify the path toward failure earlier, while keeping engineers in control of the decision.')

doc.add_heading('What not to build initially', level=1)
add_bullet(doc, 'A complete semiconductor certification platform.')
add_bullet(doc, 'A custom LSTM or Transformer before establishing simple baselines.')
add_bullet(doc, 'An automatic system that overrides the quality engineer.')
add_bullet(doc, 'A claim that leakage alone proves moisture or another physical cause.')
add_bullet(doc, 'A physical 125-degree-Celsius burn-in rig without a supervised laboratory.')
add_bullet(doc, 'A generic chatbot that does not improve testing decisions.')
add_bullet(doc, 'A dashboard containing invented aerospace accuracy figures.')

doc.add_heading('Common questions', level=1)
questions = [
    ('Does the system replace burn-in testing?', 'No. The prototype assists screening and may identify early risk. Any reduction in approved testing requires controlled industrial validation.'),
    ('Do we need Bayesian optimisation?', 'No. It may tune hyperparameters, but uncertainty-aware prediction is more valuable than simply adding Bayesian tuning.'),
    ('Can we identify exact degradation causes?', 'Only when the dataset contains trustworthy cause labels. Otherwise rank probable causes and state the evidence and uncertainty.'),
    ('Can we use synthetic data?', 'Yes for development and demonstration, provided it is physics-guided, clearly labelled and not presented as real aerospace validation.'),
    ('Why not use an LSTM?', 'Four timestamps are usually better handled by engineered features and tabular models. Deep sequence models become useful when many time points are available.'),
    ('What is our novelty?', 'The transparent combination of batch-aware anomaly detection, early drift forecasting, uncertainty and risk-sensitive recommendations.'),
]
for q, a in questions:
    add_lead(doc, q + ' ', a)

doc.add_heading('End of day self check', level=1)
checks = [
    'I can explain why a component can be dangerous while still below its absolute limit.',
    'I can distinguish an outlier score from a future-value prediction.',
    'I know why false negatives are especially serious.',
    'I can calculate change, slope and batch-relative deviation.',
    'I can explain why components and batches must remain separate across train and test sets.',
    'I can describe what Isolation Forest, XGBoost, conformal prediction and SHAP contribute.',
    'I can state which parts of our data are synthetic and what claims that prevents us from making.',
    'I can demonstrate the complete system in under three minutes.',
]
for item in checks:
    add_bullet(doc, '[ ] ' + item)

doc.add_heading('Definition of a nomination ready prototype', level=1)
doc.add_paragraph('The prototype is ready when the team can upload one dataset, reproduce a leakage-safe evaluation, show that the proposed method improves on the fixed-threshold baseline, explain at least one early warning, display uncertainty, and state the limitations of synthetic data without hesitation.')

doc.add_heading('Primary source', level=2)
doc.add_paragraph('Smart India Hackathon 2026 problem statements: https://www.sih.gov.in/sih2026PS')

doc.core_properties.title = 'SIH26170 Teammate Learning Sheet'
doc.core_properties.subject = 'One day learning guide for component burn in anomaly detection and drift prediction'
doc.core_properties.author = 'Saikarthik Ramakrishnan and team'
doc.save(OUT)
print(OUT)
